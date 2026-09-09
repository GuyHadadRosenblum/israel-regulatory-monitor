import os
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

# --- 1. DYNAMIC PATH CONFIGURATION ---
# Script location: .../regulatory_monitoring_python/src/Regulation_current_table_generator.py
SRC_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(SRC_DIR)  # Base directory: .../regulatory_monitoring_python

TABLES_BASE_DIR = os.path.join(BASE_DIR, 'data', 'Regulation current table')
OLD_TABLE_DIR = os.path.join(TABLES_BASE_DIR, 'old Regulation current table')
MONDAY_TABLE_DIR = os.path.join(TABLES_BASE_DIR, 'monday_table')
NEW_TABLE_DIR = os.path.join(TABLES_BASE_DIR, 'new Regulation current table')
MAPPING_DIR = os.path.join(TABLES_BASE_DIR, 'מיקומי שורות בבירור')

os.makedirs(NEW_TABLE_DIR, exist_ok=True)

CUTOFF_DATE = pd.to_datetime("2026-03-27")


def normalize_date_object(date_obj):
    if pd.isna(date_obj):
        return None
    return date_obj.to_pydatetime().date() if hasattr(date_obj, 'to_pydatetime') else date_obj.date()


def clean_title_string(title):
    if not title or pd.isna(title):
        return ""
    return str(title).strip().replace('"', '').replace('\\', '')


def get_first_excel_file(directory):
    if not os.path.exists(directory):
        return None
    files = [
        os.path.join(directory, f) for f in os.listdir(directory)
        if f.endswith('.xlsx') and not f.startswith('~')
    ]
    if not files:
        return None
    files.sort(key=os.path.getmtime, reverse=True)
    return files[0]


def extract_old_table_keys(file_path):
    existing_keys = set()
    df_raw = pd.read_excel(file_path, header=None)

    header_row_idx = None
    for idx, row in df_raw.iterrows():
        if 'כותרת טיוטת הרגולציה' in str(row.values):
            header_row_idx = idx
            break

    if header_row_idx is not None:
        df_old = pd.read_excel(file_path, header=header_row_idx)
        for _, row in df_old.iterrows():
            title = clean_title_string(row.get('כותרת טיוטת הרגולציה', ''))
            ministry = str(row.get('שם המשרד או הרשות', '')).strip()

            raw_date = row.get('תאריך הפנייה לרשות')
            date_val = pd.to_datetime(raw_date, errors='coerce')
            normalized_date = normalize_date_object(date_val)

            if title and title != 'nan' and normalized_date is not None:
                existing_keys.add((title, ministry, normalized_date))

    return existing_keys


def process_and_generate():
    print(f"[*] Starting Regulation Table Generation Process...")

    # --- 2. LOCATE FILES ---
    old_file_path = get_first_excel_file(OLD_TABLE_DIR)
    monday_file_path = get_first_excel_file(MONDAY_TABLE_DIR)

    if not old_file_path or not monday_file_path:
        raise FileNotFoundError('Missing historical or Monday input Excel files')

    print(f"[*] Found Old Table: {os.path.basename(old_file_path)}")
    print(f"[*] Found Monday Table: {os.path.basename(monday_file_path)}")

    # --- 2b. LOAD HISTORICAL EXCEL MAPPING DATA ---
    # File name: clash_report_old_excel_rows.xlsx
    mapping_file_xlsx = os.path.join(MAPPING_DIR, 'clash_report_old_excel_rows.xlsx')
    mapped_titles_set = set()

    if os.path.exists(mapping_file_xlsx):
        try:
            df_map = pd.read_excel(mapping_file_xlsx)
            for _, row in df_map.iterrows():
                raw_title = row.get('כותרת טיוטת הרגולציה', '')
                t_key = clean_title_string(raw_title)
                if t_key:
                    mapped_titles_set.add(t_key)
            print(
                f"[*] Successfully loaded {len(mapped_titles_set)} specific legislative filter targets from clash_report_old_excel_rows.xlsx.")
        except Exception as e:
            print(f"[!] Warning: Could not process mapping .xlsx file values: {e}")

    # --- 3. BUILD MEMORY STATE ---
    existing_keys = extract_old_table_keys(old_file_path)
    print(f"[*] Extracted {len(existing_keys)} existing records from the Old Table for cross-referencing.")

    # --- 4. PROCESS MONDAY TABLE ---
    df_monday = pd.read_excel(monday_file_path, header=2)
    df_monday.columns = df_monday.columns.str.strip()

    new_rows_to_append = []

    for index, row in df_monday.iterrows():
        # Skip rows marked as duplicates
        if str(row.get('האם כפילות', '')).strip() == 'כפילות - להתעלם':
            continue

        delayed_date = pd.to_datetime(row.get('תאריך פנייה רשמי מעוכב מעודכן'), errors='coerce')
        official_date = pd.to_datetime(row.get('תאריך פנייה רשמית'), errors='coerce')

        resolved_date_obj = delayed_date if pd.notna(delayed_date) else official_date

        if pd.isna(resolved_date_obj) or resolved_date_obj < CUTOFF_DATE:
            continue

        normalized_date = normalize_date_object(resolved_date_obj)
        title = clean_title_string(row.get('שם התקנה/חוק', ''))
        ministry = str(row.get('המשרד האחראי', '')).strip()

        row_key = (title, ministry, normalized_date)

        if row_key in existing_keys:
            continue

        portal_pub_date = row.get('תאריך פרסום באתר')
        decision_update_date = pd.to_datetime(row.get('תאריך עדכון האם הוחלט לייעץ'), errors='coerce')
        data_source = str(row.get('מקור הנתונים', '')).strip()
        ria_exists = str(row.get('האם קיים RIA/פטור', '')).strip()
        decision = str(row.get('האם הוחלט לייעץ', '')).strip()

        cond_dates_exist = pd.notna(portal_pub_date) and pd.notna(official_date)
        cond_data_source = (data_source == 'מאוחד (פורטל + פנייה)')
        cond_ria = (ria_exists == 'קיים RIA')
        cond_decision_valid = (decision != 'יש למלא')

        if cond_dates_exist and cond_data_source and cond_ria and cond_decision_valid:
            # Skip the row entirely if there is no official decision update date
            if pd.isna(decision_update_date) or str(row.get('תאריך עדכון האם הוחלט לייעץ', '')).strip() == "":
                continue

            final_decision_text = decision
            summary_text = ""

            new_rows_to_append.append({
                "title": title,
                "hyperlink": str(row.get('קישור לדף החקיקה', '')).strip(),
                "ministry": ministry,
                "date": normalized_date,
                "decision": final_decision_text,
                "update_date": normalize_date_object(decision_update_date),
                "summary": summary_text
            })
    print(f"[*] Memory check complete. Discovered {len(new_rows_to_append)} matching NEW rows to insert.")

    # --- 5. EDIT SPREADSHEET WITH OPENPYXL ---
    workbook = load_workbook(filename=old_file_path)
    sheet = workbook.active

    header_row_idx = 1
    for row_idx in range(1, min(20, sheet.max_row + 1)):
        if sheet.cell(row=row_idx, column=1).value == 'כותרת טיוטת הרגולציה':
            header_row_idx = row_idx
            break

    chronological_rows = []
    be_berur_stride_rows = []
    be_berur_append_rows = []
    original_total_rows = sheet.max_row

    for row_idx in range(header_row_idx + 1, original_total_rows + 1):
        row_data = []
        is_empty = True
        for col_idx in range(1, 7):
            cell = sheet.cell(row=row_idx, column=col_idx)
            val = cell.value
            link = cell.hyperlink.target if cell.hyperlink else None

            cell_fill_copy = None
            if cell.fill and cell.fill.fill_type:
                from openpyxl.styles import PatternFill
                cell_fill_copy = PatternFill(
                    fill_type=cell.fill.fill_type,
                    start_color=cell.fill.start_color,
                    end_color=cell.fill.end_color
                )

            row_data.append({
                'value': val,
                'link': link,
                'fill': cell_fill_copy,
                'number_format': cell.number_format
            })
            if val is not None and str(val).strip() != "":
                is_empty = False

        if not is_empty:
            title_str = clean_title_string(row_data[0]['value'])
            original_decision_val = str(row_data[3]['value']).strip()

            if original_decision_val == "בבירור":
                row_data[2]['value'] = "-"
                row_data[4]['value'] = "-"

                # Filter logic: Check if the title exists in clash_report_old_excel_rows.xlsx
                if title_str in mapped_titles_set:
                    be_berur_stride_rows.append({
                        'cells': row_data,
                        'title': title_str,
                        'original_read_row': row_idx
                    })
                else:
                    be_berur_append_rows.append({
                        'cells': row_data,
                        'title': title_str
                    })
            else:
                date_val = pd.to_datetime(row_data[2]['value'], errors='coerce')
                chronological_rows.append({
                    'cells': row_data,
                    'sort_date': date_val,
                    'title': title_str
                })

    for new_data in new_rows_to_append:
        row_data = [
            {'value': new_data["title"], 'link': new_data["hyperlink"], 'fill': None, 'number_format': None},
            {'value': new_data["ministry"], 'link': None, 'fill': None, 'number_format': None},
            {'value': new_data["date"], 'link': None, 'fill': None, 'number_format': 'dd/mm/yyyy'},
            {'value': new_data["decision"], 'link': None, 'fill': None, 'number_format': None},
            {'value': new_data["update_date"], 'link': None, 'fill': None, 'number_format': 'dd/mm/yyyy'},
            {'value': new_data["summary"], 'link': None, 'fill': None, 'number_format': None}
        ]
        date_val = pd.to_datetime(new_data["date"], errors='coerce')
        chronological_rows.append({
            'cells': row_data,
            'sort_date': date_val,
            'title': new_data["title"]
        })

    # Sort chronological data entries natively by Date
    chronological_rows.sort(key=lambda x: x['sort_date'] if pd.notna(x['sort_date']) else pd.Timestamp.min)
    final_data_rows = [r['cells'] for r in chronological_rows]

    # --- 6. UNIFORM DISTRIBUTION FOR SPECIFIC MAPPED 'בבירור' ENTRIES ---
    print("\n=================== UNIFORM SPREAD PLACEMENT LOG ===================")
    be_berur_final_log = []

    if be_berur_stride_rows and final_data_rows:
        # Stride calculation based strictly on the count of filtered mapped items
        stride = max(1, len(final_data_rows) // len(be_berur_stride_rows))

        for i, b_item in enumerate(be_berur_stride_rows):
            target_list_idx = (i * stride) + (stride // 2)
            target_list_idx = max(0, min(target_list_idx, len(final_data_rows)))

            final_data_rows.insert(target_list_idx, b_item['cells'])

            expected_excel_row = header_row_idx + 1 + target_list_idx
            be_berur_final_log.append({
                'title': b_item['title'],
                'type': 'Uniform Spread',
                'new_row': expected_excel_row
            })
            print(
                f"[SPREAD] 'בבירור' (Mapped) -> Placed at Row Index: {target_list_idx} | Title: {b_item['title'][:40]}...")
    else:
        for b_item in be_berur_stride_rows:
            final_data_rows.append(b_item['cells'])
            be_berur_final_log.append({'title': b_item['title'], 'type': 'Append (Fallback)',
                                       'new_row': header_row_idx + len(final_data_rows)})

    # Append standard 'בבירור' rows that were not found in the configuration mapping file
    for b_item in be_berur_append_rows:
        final_data_rows.append(b_item['cells'])
        expected_excel_row = header_row_idx + len(final_data_rows)
        be_berur_final_log.append({
            'title': b_item['title'],
            'type': 'Bottom Append (Unmapped)',
            'new_row': expected_excel_row
        })
        print(
            f"[APPEND] 'בבירור' (Unmapped) -> Appended to bottom at row: {expected_excel_row} | Title: {b_item['title'][:40]}...")

    # Styles and layout configuration
    font_standard = Font(name="Calibri", size=11)
    font_hyperlink = Font(name="Calibri", size=11, color="0066CC", underline="single")
    alignment_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    sheet.row_dimensions[header_row_idx].height = 28

    print("\n=================== FINAL EXCEL WRITING LOG ===================")
    current_row_idx = header_row_idx + 1
    for row_cells in final_data_rows:
        sheet.row_dimensions[current_row_idx].height = 24

        decision_txt = str(row_cells[3]['value']).strip()
        title_txt = str(row_cells[0]['value']).strip()

        if decision_txt == "בבירור":
            print(f"[WRITE] Row {current_row_idx:03d} -> Written 'בבירור' row: \"{title_txt[:40]}...\"")

        for col_idx in range(1, 7):
            cell = sheet.cell(row=current_row_idx, column=col_idx)
            cell_data = row_cells[col_idx - 1]

            cell.value = cell_data['value']
            cell.alignment = alignment_center
            cell.border = thin_border

            if cell_data.get('fill') is not None and cell_data['fill'].fill_type is not None:
                cell.fill = cell_data['fill']
            else:
                from openpyxl.styles import PatternFill
                cell.fill = PatternFill(fill_type=None)

            if str(cell.value).strip() == "-":
                cell.number_format = '@'

            if cell_data['link'] and str(cell_data['link']).startswith("http"):
                cell.hyperlink = cell_data['link']
                cell.font = font_hyperlink
            elif col_idx == 1 and current_row_idx >= 6:
                cell.hyperlink = None
                cell.font = font_hyperlink
            else:
                cell.hyperlink = None
                cell.font = font_standard

        current_row_idx += 1
    print("================================================================\n")

    # Print comparative placement summary
    print("=================== SUMMARY: DISTRIBUTION PLACEMENT RESULTS ===================")
    for item in be_berur_final_log:
        print(f"Type: {item['type']:<26} | New Written Row: {item['new_row']:<3} | Title: {item['title'][:45]}")
    print("===============================================================================\n")

    for tbl_name in list(sheet.tables.keys()):
        del sheet.tables[tbl_name]

    table_ref = f"A{header_row_idx}:F{current_row_idx - 1}"
    tab = Table(displayName="RegulationUpdatesTable", ref=table_ref)

    style = TableStyleInfo(
        name="TableStyleLight1", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    tab.tableStyleInfo = style
    sheet.add_table(tab)

    for r_idx in range(header_row_idx + 1, current_row_idx):
        for c_idx in [3, 4, 5]:
            target_cell = sheet.cell(row=r_idx, column=c_idx)
            if isinstance(target_cell.value, (datetime, pd.Timestamp, type(datetime.now().date()))):
                target_cell.number_format = 'dd/mm/yyyy'

    sheet['F2'] = datetime.now().date()
    sheet['F2'].number_format = 'dd/mm/yyyy'

    timestamp_str = datetime.now().strftime("%Y-%m-%d_%H-%M")
    base_filename = f"Regulation_Tracking_Update_{timestamp_str}"
    new_filename = f"{base_filename}.xlsx"
    new_file_path = os.path.join(NEW_TABLE_DIR, new_filename)

    counter = 1
    while os.path.exists(new_file_path):
        try:
            with open(new_file_path, 'r+'):
                break
        except IOError:
            new_filename = f"{base_filename}_copy{counter}.xlsx"
            new_file_path = os.path.join(NEW_TABLE_DIR, new_filename)
            counter += 1

    workbook.save(filename=new_file_path)
    print(f"[V] SUCCESS! Sorted, styled, and saved at:\n    {new_file_path}")

    old_dir_mirror_path = os.path.join(OLD_TABLE_DIR, new_filename)
    workbook.save(filename=old_dir_mirror_path)
    print(f"[V] Successfully added latest state version to history in source directory:\n    {old_dir_mirror_path}")


if __name__ == "__main__":
    process_and_generate()
